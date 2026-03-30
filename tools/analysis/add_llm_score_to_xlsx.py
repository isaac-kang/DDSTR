"""Add LLM confidence score row to error_details_PL.xlsx.

Reads the xlsx (8-row blocks per sample), looks up LLM scores from cache,
inserts an 'LLM' row (with σ(r_llm)) below the 'PL' row, and saves to a new file.
Preserves all cell formatting (font colors, fills, etc.).

Usage:
    python tools/analysis/add_llm_score_to_xlsx.py
    python tools/analysis/add_llm_score_to_xlsx.py --input tools/analysis/error_details_PL.xlsx --llm_cache tools/analysis/llm_scores_cache_Qwen-Qwen3-32B-AWQ.json
"""

import argparse
import json
import os
from copy import copy
from pathlib import Path

import numpy as np
import openpyxl


def sigmoid(x):
    return 1.0 / (1.0 + np.exp(-np.clip(x, -500, 500)))


def copy_cell(src, dst):
    """Copy value and formatting from src cell to dst cell."""
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', default='tools/analysis/error_details_PL.xlsx')
    parser.add_argument('--llm_cache', default='tools/analysis/llm_scores_cache_Qwen-Qwen3-32B-AWQ.json')
    parser.add_argument('--output', default=None,
                        help='Output xlsx path (default: input with _llm suffix)')
    args = parser.parse_args()

    if args.output is None:
        stem = Path(args.input).stem
        args.output = str(Path(args.input).parent / f'{stem}_llm.xlsx')

    # Load LLM scores
    with open(args.llm_cache, 'r') as f:
        cache = json.load(f)
    scores_list = cache if isinstance(cache, list) else cache['scores']
    llm_lookup = {(r['dataset_name'], r['image_index']): r['r_llm'] for r in scores_list}
    llm_model = cache.get('llm_model', 'unknown') if isinstance(cache, dict) else 'unknown'
    print(f'Loaded LLM scores: {len(llm_lookup)} entries (model: {llm_model})')

    # Read source xlsx
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active
    max_row = ws.max_row
    max_col = ws.max_column

    # Collect source rows as cell objects (1-indexed)
    src_rows = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        src_rows.append(row)

    # Build output: list of (type, data) where type is 'copy' (src row index) or 'llm' (score value)
    output_plan = []
    i = 0
    count = 0
    while i < len(src_rows):
        first_cell = src_rows[i][0]
        # Detect block start: col A has dataset_name (non-None string)
        if first_cell.value is not None and isinstance(first_cell.value, str):
            dataset_name = first_cell.value
            image_index = src_rows[i][1].value

            # Copy the 7 data rows
            block_end = min(i + 7, len(src_rows))
            for j in range(i, block_end):
                output_plan.append(('copy', j))

            # Look up LLM score and insert row
            r_llm = llm_lookup.get((dataset_name, image_index), None)
            if r_llm is not None:
                conf_llm = float(sigmoid(r_llm))
                output_plan.append(('llm', round(conf_llm, 4)))
                count += 1
            else:
                output_plan.append(('llm', 'N/A'))

            # Blank separator
            output_plan.append(('blank', None))
            # Skip original block (7 data rows + 1 blank)
            i = block_end + 1
        else:
            output_plan.append(('copy', i))
            i += 1

    # Write new xlsx, preserving formatting
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = ws.title

    out_row = 1
    for action, data in output_plan:
        if action == 'copy':
            src_row = src_rows[data]
            for col_idx, src_cell in enumerate(src_row, start=1):
                dst_cell = ws_out.cell(row=out_row, column=col_idx)
                copy_cell(src_cell, dst_cell)
        elif action == 'llm':
            ws_out.cell(row=out_row, column=2, value='LLM')
            ws_out.cell(row=out_row, column=3, value=data)
        # 'blank' -> leave row empty
        out_row += 1

    # Copy column widths
    for col_letter, dim in ws.column_dimensions.items():
        ws_out.column_dimensions[col_letter].width = dim.width

    os.makedirs(os.path.dirname(args.output) or '.', exist_ok=True)
    wb_out.save(args.output)
    print(f'Added LLM scores to {count} samples')
    print(f'Saved to {args.output}')

    wb.close()


if __name__ == '__main__':
    main()
